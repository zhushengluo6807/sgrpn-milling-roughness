# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


OUT = Path("切削实验") / "物理模型神经网络最终调整方案.docx"


def set_cell_text(cell, text):
    cell.text = text
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)


def set_doc_style(doc):
    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.bold = True


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        rest = text[len(bold_prefix):]
        if rest:
            p.add_run(rest)
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        for run in p.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        for run in p.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def main():
    doc = Document()
    set_doc_style(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("物理模型神经网络最终调整方案")
    r.bold = True
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("面向铣削表面粗糙度预测的机理化物理基准、振动残差补偿与门控融合框架")
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(11)

    doc.add_heading("一、总体定位", level=1)
    add_para(doc, "本文方案建议定位为：面向变工况铣削表面粗糙度预测的机理化物理基准引导残差学习方法。")
    add_para(doc, "核心思想不是使用复杂黑箱网络直接拟合粗糙度，而是先建立比传统几何公式更贴近槽铣过程的物理基准，再利用主轴振动信号学习实际加工中的动态偏差。")
    add_para(doc, "最终模型主线为：机理化表面生成物理基准 + 振动特征残差网络 + 多因素门控残差补偿。")
    add_para(doc, "当前数据基础为：共设计 224 组实验，最终获得 212 组有效稳定切削信号，12 组缺失，总缺失率约 5.36%。若每组有效信号划分为 3 个非重叠片段，可形成 636 个训练片段，但独立实验样本数仍为 212。")

    doc.add_heading("二、为什么不再以简单几何公式作为最终物理模型", level=1)
    add_para(doc, "传统几何粗糙度公式 Ra_geo = fz² / (32Reff) 只描述理想刀具几何与进给运动形成的残留高度，适合作为对照基线，但不适合作为最终物理基准。")
    add_bullets(doc, [
        "该公式只显式体现每齿进给量和等效刀尖半径，无法体现轴向切深、切削载荷、刀具跳动和刀具让刀。",
        "槽铣属于全径向切削，切削负载较大，实际表面生成会受到多齿切削不均和刀具-主轴系统变形影响。",
        "如果继续只使用该公式，后续残差网络需要补偿过多物理模型未覆盖的偏差，物理分支容易变弱。"
    ])
    add_para(doc, "因此，本文建议将该公式作为“基础几何基线模型”，而将最终物理基准升级为机理化表面生成模型。")

    doc.add_heading("三、最终物理基准：考虑径向跳动与静态让刀的槽铣表面生成模型", level=1)
    add_para(doc, "最终物理基准建议命名为：考虑径向跳动与静态让刀的槽铣表面生成物理基准模型。")
    add_para(doc, "该模型不是在 Ra_geo 上加入经验 β 修正项，而是基于刀齿运动轨迹、刀具径向跳动和静态让刀重新生成加工表面轮廓，再由轮廓计算 Ra_phy。")

    doc.add_heading("3.1 模型输入", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["符号", "含义", "来源建议"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h)
        shade_cell(table.rows[0].cells[i], "D9EAF7")
    rows = [
        ("fz", "每齿进给量", "实验记录表"),
        ("n", "主轴转速", "实验记录表"),
        ("ap", "轴向切深", "实验记录表"),
        ("R", "刀具半径", "刀具规格"),
        ("Z", "刀齿数", "本文为 3 齿"),
        ("e", "刀具径向跳动量", "未实测时作为不确定结构参数做敏感性分析"),
        ("k_tool", "刀具-主轴系统等效刚度", "简化梁模型、文献范围或敏感性分析"),
        ("Kc", "切削力系数", "文献或材料经验值，避免由粗糙度标签自由拟合")
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)

    doc.add_heading("3.2 建模流程", level=2)
    add_numbered(doc, [
        "建立多齿立铣刀运动轨迹，根据主轴转速、刀齿数和进给速度描述每个刀齿的相对运动。",
        "引入径向跳动量 e，使不同刀齿具有不同的等效旋转半径，从而描述齿间切削轨迹偏置。",
        "根据每齿进给量和轴向切深估算未变形切屑厚度与名义切削力。",
        "根据等效刚度 k_tool 估算静态让刀量，使刀具实际轨迹偏离理想轨迹。",
        "将考虑跳动和让刀后的刀齿轨迹映射到工件表面，生成理论表面轮廓。",
        "对生成轮廓计算算术平均粗糙度，得到机理化物理基准 Ra_phy。"
    ])

    doc.add_heading("3.3 模型参考依据", level=2)
    add_bullets(doc, [
        "铣削表面生成理论：通过刀具运动轨迹和材料去除包络生成加工表面轮廓，是 surface generation / surface topography prediction 的常见思想。",
        "刀具径向跳动机理：实际多齿铣削中各刀齿旋转半径并不完全一致，会导致齿间切削厚度不均和表面残留轮廓变化。",
        "静态让刀机理：槽铣切削负载较大，轴向切深增加会提高切削力，引起刀具-主轴系统弹性变形，使实际刀具轨迹偏离理想几何轨迹。",
        "铣削力模型基础：可参考经典铣削力学中未变形切屑厚度与切削力系数建模思想。"
    ])
    add_para(doc, "论文中应避免声称建立了高保真切削仿真模型，更稳妥的表述是：构建具有明确切削机理约束的物理粗糙度基准，用于为神经网络提供比简单几何公式更可靠的先验锚点。")

    doc.add_heading("四、无实测径向跳动数据时的处理", level=1)
    add_para(doc, "由于当前实验未单独测量径向跳动，径向跳动 e 不应写成实测输入，也不建议通过粗糙度标签反演为拟合参数。")
    add_para(doc, "推荐写法：将 e 视为机床-刀柄-刀具装夹系统中的不确定结构参数，并在合理工程范围内开展敏感性分析。")
    add_para(doc, "建议取值范围：e = 0、5、10、15、20 μm。其中 0 μm 表示理想无跳动，5-10 μm 表示较常见装夹精度，15-20 μm 表示跳动较明显工况。")

    doc.add_heading("五、敏感性分析方案", level=1)
    add_para(doc, "敏感性分析分为两层：物理基准敏感性分析和最终预测敏感性分析。")

    doc.add_heading("5.1 物理基准敏感性分析", level=2)
    add_bullets(doc, [
        "固定其他实验参数，只改变径向跳动 e。",
        "分别计算不同 e 下的 Ra_phy(e)。",
        "统计 Ra_phy(e) 相对于 e=0 μm 的变化百分比。",
        "绘制 e-Ra_phy 曲线或 e-物理基准误差曲线。",
    ])
    add_para(doc, "可使用如下指标：S(e)=mean(|Ra_phy(e)-Ra_phy(0)| / Ra_phy(0))。该指标用于说明径向跳动对物理基准是否具有显著影响。")

    doc.add_heading("5.2 最终模型敏感性分析", level=2)
    add_bullets(doc, [
        "对 e = 0、5、10、15、20 μm 分别训练和测试最终模型。",
        "记录 MAE、RMSE、R²、MAPE 等指标。",
        "若不同 e 下性能差异较小，说明门控残差模型对径向跳动不确定性具有鲁棒性。",
        "若某个 e 下性能最优，只能说明该量级与实验观测更一致，不能解释为真实径向跳动测量值。"
    ])

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["径向跳动 e(μm)", "物理基准 MAE", "最终模型 MAE", "RMSE", "R²"]):
        set_cell_text(table.rows[0].cells[i], h)
        shade_cell(table.rows[0].cells[i], "D9EAD3")
    for e in ["0", "5", "10", "15", "20"]:
        cells = table.add_row().cells
        for i, val in enumerate([e, "", "", "", ""]):
            set_cell_text(cells[i], val)

    doc.add_heading("六、残差网络设计", level=1)
    add_para(doc, "物理模型负责输出 Ra_phy，残差网络只学习物理模型难以覆盖的动态偏差。")
    add_para(doc, "推荐最终形式：Ra_pred = Ra_phy + α·ΔRa_net。")
    add_para(doc, "该形式的解释是：Ra_phy 为机理化物理基准，ΔRa_net 为由振动信号和工况特征学习到的动态偏差，α 控制动态残差补偿强度。")

    doc.add_heading("6.1 特征输入", level=2)
    add_bullets(doc, [
        "切削参数：n、fz、ap。",
        "物理基准：Ra_phy。",
        "人工时域特征：RMS、峰峰值、标准差、峭度、偏度、峰值因子等。",
        "人工频域特征：主轴转频能量、刀齿通过频率能量、倍频能量、高频能量占比、频谱质心等。",
        "轻量 1D-CNN 特征：从三轴振动信号中提取 16-32 维深度特征。"
    ])

    doc.add_heading("6.2 网络体量", level=2)
    add_para(doc, "当独立实验数据约 200 条以上、切分后训练片段约 600 条时，建议残差网络采用如下体量：")
    add_para(doc, "Linear(input_dim, 64) → ReLU → Dropout(0.2) → Linear(64, 32) → ReLU → Linear(32, 1)")
    add_para(doc, "如果最终独立实验数量不足 200 条，则建议缩小为：")
    add_para(doc, "Linear(input_dim, 32) → ReLU → Linear(32, 16) → ReLU → Linear(16, 1)")

    doc.add_heading("七、轻量 1D-CNN 特征提取", level=1)
    add_para(doc, "在当前 212 组独立实验、约 636 个训练片段的数据规模下，可以将轻量 1D-CNN 作为正式特征分支，而不仅是辅助对比模块。")
    add_para(doc, "仍不建议使用大型端到端 CNN、STFT-2D-CNN、CNN-LSTM 或 Transformer。推荐采用人工时频特征 + 轻量 1D-CNN 的双特征路线。")
    add_para(doc, "轻量 1D-CNN 推荐结构：Conv1D(16, kernel=32, stride=4) → ReLU → Conv1D(32, kernel=16, stride=2) → ReLU → GlobalAveragePooling → 16-32维深度特征。")
    add_para(doc, "这样既能保留振动信号局部波形信息，又不会在小样本条件下引入过多参数。")

    doc.add_heading("八、门控拓展方案", level=1)
    add_para(doc, "门控不再只依赖主轴转速，而应使用多因素输入，以判断当前工况下动态残差补偿的必要程度。")
    add_bullets(doc, [
        "门控输入：n、fz、ap、Ra_phy、RMS、刀齿通过频率能量、高频能量占比。",
        "门控结构：Linear(gate_input_dim, 8) → ReLU → Linear(8, 1) → Sigmoid。",
        "门控输出：α ∈ [0, 1]。",
        "最终预测：Ra_pred = Ra_phy + α·ΔRa_net。"
    ])
    add_para(doc, "物理解释：稳定切削、振动较小时，α 较小，说明物理基准已能解释主要粗糙度来源；大切深、高振动或频率能量异常增强时，α 较大，说明需要更多动态残差补偿。")

    doc.add_heading("九、数据使用与划分", level=1)
    add_para(doc, "当前共设计 224 组实验，其中 212 组获得有效稳定切削信号，12 组缺失。将每条有效稳定切削信号划分为 3 个非重叠片段后，可形成 636 个训练片段。")
    add_para(doc, "必须强调：切分片段增加的是训练片段数，不是独立实验数；论文中应同时报告独立实验数 212 与训练片段数 636。")
    add_bullets(doc, [
        "同一执行顺序切出的三个片段必须同时属于训练集、验证集或测试集。",
        "不得让同一组加工参数的不同片段同时出现在训练集和测试集中。",
        "建议采用 GroupKFold 或按执行顺序分组的数据划分策略，所有模型对比与消融实验均应遵守该划分方式。",
        "缺失数据不参与训练；质量可疑数据可作为扩展数据集或敏感性分析使用，主结果优先基于高质量数据集报告。"
    ])

    doc.add_heading("9.1 缺失数据分布与风险说明", level=2)
    add_para(doc, "缺失数据总体不严重，但并非完全均匀分布。v3 缺失 3 条，执行顺序为 26、79、105；v4 缺失 9 条，执行顺序为 1、2、27、28、40、66、79、92、105。")
    add_bullets(doc, [
        "按切深看：v4 的 0.5 mm 和 1.5 mm 切深缺失相对较多，分别缺失 4 条；2.0 mm 切深未缺失。",
        "按转速看：v4 的 4500 rpm 和 5500 rpm 各缺失 3 条，低/中低转速缺失相对较多。",
        "按进给看：v4 的 0.06 mm/z 缺失 5 条，是最明显的局部偏向。",
        "按块号/层号看：v4 层 1 缺失 6 条、层 2 缺失 3 条，存在轻微偏向，但没有某一整块或整层完全缺失。"
    ])
    add_para(doc, "论文中建议表述为：缺失样本主要由信号截取失败或有效切削段质量不足导致。总体缺失率较低，且各主要转速、进给和切深水平仍保留有效样本，因此未破坏整体参数空间覆盖。但 v4 数据中低进给及部分切深水平存在一定缺失偏向，后续模型训练与评价中采用分组交叉验证和敏感性分析降低其影响。")

    doc.add_heading("十、实验验证设计", level=1)
    add_para(doc, "为支撑论文结论，建议至少完成以下对比和消融实验。")
    add_bullets(doc, [
        "对比模型：传统几何模型、机理化物理基准、SVR、Random Forest/XGBoost、MLP、纯数据驱动模型、本文模型。",
        "消融实验：无物理基准、仅 Ra_geo、Ra_phy、Ra_phy+残差、Ra_phy+门控残差、加入/不加入轻量 1D-CNN。",
        "敏感性实验：不同径向跳动 e 下的物理基准和最终预测性能。",
        "数据质量实验：高质量数据集 vs 高质量数据集+可疑数据。",
        "缺失偏向分析：报告 v4 中 0.06 mm/z、0.5 mm、1.5 mm、4500 rpm 和 5500 rpm 的局部缺失情况，并说明模型验证采用分组划分以降低偏向影响。",
        "可解释性分析：门控 α 随转速、切深、振动 RMS、刀齿通过频率能量变化的趋势。"
    ])

    doc.add_heading("十一、论文高度判断", level=1)
    add_para(doc, "经过上述调整后，论文的核心卖点将从简单的“神经网络预测粗糙度”提升为“机理化表面生成物理基准 + 振动残差补偿 + 多因素门控融合”。")
    add_bullets(doc, [
        "硕士论文核心章节：较扎实。",
        "中文核心或 EI 工程应用类论文：希望较大。",
        "普通 SCI 工程应用类论文：有机会。",
        "较好 SCI 三区或二区边缘：取决于新增数据质量、对比实验完整性和物理模型验证充分性。",
        "机械或智能制造顶刊：仍需要更多实测物理量，例如切削力、刀具跳动、刀具磨损或温度数据支撑。"
    ])

    doc.add_heading("十二、最终推荐方案一句话概括", level=1)
    add_para(doc, "本文建议采用考虑刀具径向跳动与静态让刀的槽铣表面生成物理基准模型，结合主轴振动时频特征与轻量 1D-CNN 特征，通过小型残差网络学习动态加工偏差，并利用多因素门控机制自适应调节残差补偿强度；在 212 组独立实验和 636 个训练片段的基础上，通过分组验证、缺失偏向说明、敏感性分析和消融实验，实现兼具预测精度与物理可解释性的铣削表面粗糙度预测。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
