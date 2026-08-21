import csv
from pathlib import Path

import numpy as np
import openpyxl


FS_EXPECTED = 25600.0
SLOT_LENGTH_MM = 120.0
STABLE_LENGTH_MM = 100.0
EDGE_TRIM_MM = (SLOT_LENGTH_MM - STABLE_LENGTH_MM) / 2.0


def find_data_dir(root: Path) -> Path:
    for path in root.iterdir():
        if path.is_dir() and (path / "1.txt").exists():
            return path
    raise FileNotFoundError("Could not find the experiment data directory.")


def find_record_file(root: Path) -> Path:
    for path in root.glob("*.xlsx"):
        if path.name.startswith("实验记录"):
            return path
    raise FileNotFoundError("Could not find 实验记录表.xlsx.")


def read_metadata(root: Path):
    wb = openpyxl.load_workbook(find_record_file(root), data_only=True)
    ws = wb.worksheets[0]
    rows = {}
    for r in range(4, ws.max_row + 1):
        seq = ws.cell(r, 1).value
        if not isinstance(seq, (int, float)):
            continue
        rows[int(seq)] = {
            "seq": int(seq),
            "condition_id": ws.cell(r, 2).value,
            "block": ws.cell(r, 3).value,
            "layer": ws.cell(r, 5).value,
            "slot": str(ws.cell(r, 6).value or "").replace("\n", "/"),
            "rpm": float(ws.cell(r, 7).value),
            "fz": float(ws.cell(r, 8).value),
            "ap": float(ws.cell(r, 9).value),
            "vf": float(ws.cell(r, 10).value),
            "ra_mean": ws.cell(r, 15).value,
            "note": ws.cell(r, 20).value,
        }
    return rows


def block_rms(vector: np.ndarray, block: int):
    n = len(vector) // block
    trimmed = vector[: n * block]
    return np.sqrt(np.mean(trimmed.reshape(n, block) ** 2, axis=1))


def moving_average(y: np.ndarray, width: int):
    if width <= 1:
        return y
    return np.convolve(y, np.ones(width) / width, mode="same")


def robust_threshold(envelope: np.ndarray):
    lower = envelope[envelope <= np.percentile(envelope, 35)]
    base = float(np.median(lower))
    mad = float(np.median(np.abs(lower - base)))
    sigma = 1.4826 * mad + 1e-12
    spread = float(np.percentile(envelope, 95) - base)
    return base, max(base + 6.0 * sigma, base + 0.20 * spread)


def first_sustained_crossing(times, envelope, threshold, min_duration=0.12):
    mask = envelope > threshold
    min_blocks = max(1, int(round(min_duration / np.median(np.diff(times)))))
    i = 0
    while i < len(mask):
        if not mask[i]:
            i += 1
            continue
        j = i
        while j < len(mask) and mask[j]:
            j += 1
        if j - i >= min_blocks:
            return float(times[i])
        i = j
    return None


def strongest_pulse_near(times, vector, center, window):
    if center is None:
        return None, None
    diff = np.abs(np.diff(vector))
    diff_t = times[:-1]
    mask = (diff_t >= center - window) & (diff_t <= center + window)
    if not np.any(mask):
        return None, None
    local = diff[mask]
    local_t = diff_t[mask]
    idx = int(np.argmax(local))
    return float(local_t[idx]), float(local[idx])


def jump_stats(x: np.ndarray):
    d = np.diff(x, axis=0)
    med = np.median(d, axis=0)
    mad = np.median(np.abs(d - med), axis=0) + 1e-12
    z = np.abs(d - med) / (1.4826 * mad)
    return {
        "jump_gt20": int(np.sum(z > 20)),
        "jump_gt50": int(np.sum(z > 50)),
        "max_robust_jump_z": float(np.max(z)),
        "max_abs_g": float(np.max(np.abs(x))),
    }


def detect_one(path: Path, meta: dict):
    arr = np.loadtxt(str(path), skiprows=2)
    t = arr[:, 0]
    x = arr[:, 1:4]
    dt = float(np.median(np.diff(t[: min(2000, len(t))])))
    fs = 1.0 / dt
    vector = np.sqrt(np.sum(x * x, axis=1))

    block = max(1, int(round(0.05 * fs)))
    env = block_rms(vector, block)
    env = moving_average(env, 3)
    env_t = (np.arange(len(env)) + 0.5) * block / fs
    base, threshold = robust_threshold(env)

    entry = first_sustained_crossing(env_t, env, threshold)
    entry_confidence = "high"
    if entry is None:
        idx = int(np.argmax(env))
        entry = float(env_t[idx])
        entry_confidence = "low"

    vf = float(meta["vf"])
    cut_duration = SLOT_LENGTH_MM / vf * 60.0
    stable_duration = STABLE_LENGTH_MM / vf * 60.0
    theory_exit = entry + cut_duration
    stable_start = entry + EDGE_TRIM_MM / vf * 60.0
    stable_end = stable_start + stable_duration

    pulse_window = max(1.0, min(2.5, 0.35 * cut_duration))
    observed_exit, observed_exit_jump = strongest_pulse_near(t, vector, theory_exit, pulse_window)
    if observed_exit_jump is not None:
        global_diff = np.abs(np.diff(vector))
        high_pulse = observed_exit_jump >= np.percentile(global_diff, 99.9)
    else:
        high_pulse = False
    if not high_pulse:
        observed_exit = None

    stats = jump_stats(x)
    captured_duration = float(t[-1])
    stable_coverage = max(0.0, min(captured_duration, stable_end) - stable_start) / stable_duration

    flags = []
    if entry_confidence == "low":
        flags.append("entry_low_confidence")
    if theory_exit > captured_duration:
        flags.append("theoretical_exit_after_record_end")
    if stable_end > captured_duration:
        flags.append("stable_segment_incomplete")
    if observed_exit is None:
        flags.append("observed_exit_not_clear")
    else:
        diff = abs(observed_exit - theory_exit)
        if diff > 1.0:
            flags.append("observed_exit_far_from_theory")
    if stats["max_abs_g"] > 1.0:
        flags.append("extreme_spike_gt1g")
    elif stats["max_abs_g"] > 0.3:
        flags.append("large_spike_gt0.3g")
    if stats["jump_gt50"] > 10:
        flags.append("many_abrupt_jumps")

    preprocess_start = stable_start
    preprocess_end = min(stable_end, captured_duration)
    if "entry_low_confidence" in flags or stable_coverage < 0.8:
        severity = "hard_review"
        action = "人工复核进刀/稳定段；不建议直接进入训练"
    elif "stable_segment_incomplete" in flags:
        severity = "partial"
        action = "稳定段末端未完整采到；可保留并按 coverage 降权"
    elif "extreme_spike_gt1g" in flags or "many_abrupt_jumps" in flags:
        severity = "repair"
        action = "先做跳变点检测和插值修复，再截取稳定段"
    elif "observed_exit_far_from_theory" in flags or "observed_exit_not_clear" in flags:
        severity = "review_exit"
        action = "出刀冲击不清楚或与理论偏差较大；预处理优先使用理论稳定段"
    else:
        severity = "ok"
        action = "可按理论稳定段直接截取"

    return {
        "seq": meta["seq"],
        "file": path.name,
        "condition_id": meta["condition_id"],
        "block": meta["block"],
        "layer": meta["layer"],
        "slot": meta["slot"],
        "rpm": meta["rpm"],
        "fz": meta["fz"],
        "ap": meta["ap"],
        "vf_mm_min": vf,
        "record_duration_s": captured_duration,
        "fs_hz": fs,
        "entry_s": entry,
        "entry_confidence": entry_confidence,
        "theoretical_exit_120mm_s": theory_exit,
        "observed_exit_pulse_s": observed_exit,
        "observed_vs_theory_delta_s": None if observed_exit is None else observed_exit - theory_exit,
        "stable_100mm_start_s": stable_start,
        "stable_100mm_end_s": stable_end,
        "preprocess_start_s": preprocess_start,
        "preprocess_end_s": preprocess_end,
        "stable_coverage": stable_coverage,
        "envelope_baseline": base,
        "envelope_threshold": threshold,
        **stats,
        "flags": ";".join(flags) if flags else "ok",
        "severity": severity,
        "recommended_action": action,
        "record_note": meta.get("note") or "",
    }


def main():
    root = Path.cwd()
    data_dir = find_data_dir(root)
    metadata = read_metadata(root)
    out_dir = root / "outputs" / "preprocessing"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_csv = out_dir / "entry_exit_detection_summary.csv"

    rows = []
    for seq in sorted(metadata):
        path = data_dir / f"{seq}.txt"
        if not path.exists():
            row = {**metadata[seq], "file": path.name, "flags": "missing_txt_file"}
        else:
            row = detect_one(path, metadata[seq])
        rows.append(row)
        print(f"{seq:03d} {row.get('flags', '')}")

    fieldnames = [
        "seq",
        "file",
        "condition_id",
        "block",
        "layer",
        "slot",
        "rpm",
        "fz",
        "ap",
        "vf_mm_min",
        "record_duration_s",
        "fs_hz",
        "entry_s",
        "entry_confidence",
        "theoretical_exit_120mm_s",
        "observed_exit_pulse_s",
        "observed_vs_theory_delta_s",
        "stable_100mm_start_s",
        "stable_100mm_end_s",
        "preprocess_start_s",
        "preprocess_end_s",
        "stable_coverage",
        "max_abs_g",
        "jump_gt20",
        "jump_gt50",
        "max_robust_jump_z",
        "envelope_baseline",
        "envelope_threshold",
        "flags",
        "severity",
        "recommended_action",
        "record_note",
    ]
    with out_csv.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    total = len(rows)
    ok = sum(1 for r in rows if r["flags"] == "ok")
    print(f"OUTPUT {out_csv}")
    print(f"TOTAL {total} OK {ok} FLAGGED {total-ok}")


if __name__ == "__main__":
    main()
