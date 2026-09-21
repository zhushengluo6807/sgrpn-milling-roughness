from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "切削实验" / "振动信号预处理理论依据与现实依据.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_cm):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_para(doc, text="", style=None, bold_label=None):
    p = doc.add_paragraph(style=style)
    if bold_label:
        r = p.add_run(bold_label)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "SimHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_step(doc, title, role, theory, reality, note):
    doc.add_heading(title, level=2)
    add_para(doc, role, bold_label="作用：")
    add_para(doc, theory, bold_label="理论依据：")
    add_para(doc, reality, bold_label="本数据现实依据：")
    add_para(doc, note, bold_label="参数与使用说明：")


def main():
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("振动信号预处理理论依据与现实依据")
    run.font.name = "SimHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(20)
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("适用于铣削稳定切削段振动信号建模前处理").italic = True

    doc.add_heading("一、数据与预处理配置", level=1)
    add_para(
        doc,
        "本批数据为已截取出的单次加工有效切削信号，共 212 条有效独立实验信号；预处理后再切分得到 586 个训练片段。采样率为 25600 Hz，奈奎斯特频率为 12800 Hz。",
    )
    add_para(
        doc,
        "当前预处理流程为：中位数去直流偏置、0.05 s 滑动窗口 RMS/峰峰值异常检测、MAD 稳健阈值判断、异常窗口前后扩展 0.2 s 剔除、20-10000 Hz FFT 带通滤波、重建连续时间轴，然后再按时长规则切分。",
    )
    add_para(
        doc,
        "转速范围约为 3500、4500、5500 rpm；若刀具为三刃，则主轴转频约为 58.3、75.0、91.7 Hz，刀齿通过频率约为 175、225、275 Hz。",
    )

    doc.add_heading("二、各项预处理依据", level=1)

    add_step(
        doc,
        "1. 中位数去直流偏置",
        "消除每条信号各通道的零点偏移，使振动围绕零基线波动，提升不同实验之间的可比性。",
        "振动采集中常见传感器零漂、安装预紧差异和采集系统偏置。直流偏置通常不代表切削动力学响应，去除偏置可以避免均值类、能量类和 CNN 输入受到基线位置影响。采用中位数而不是均值，是因为中位数对局部冲击和异常峰值更稳健。",
        "你的数据中不同文件来源、不同块号/层号和不同截取区间可能存在基线差异。用每通道中位数去偏置，可以在不改变振动波动形态的前提下削弱采集条件差异。",
        "处理方式为每个通道分别减去该通道中位数。该步骤属于信号级预处理，应在切分片段之前完成。",
    )

    add_step(
        doc,
        "2. 滑动窗口 RMS 与峰峰值检测",
        "用局部窗口判断振动强度和瞬时冲击是否异常，主要服务于异常段识别。",
        "RMS 反映一段时间内的振动能量，是加工状态监测和刀具状态监测中的常用时域指标；峰峰值对短时冲击敏感，适合发现切入切出、抬刀、主轴停转、碰撞或采集异常造成的尖峰波动。两者结合可以同时覆盖“能量持续偏大”和“瞬时幅值突变”两类异常。",
        "采样率为 25600 Hz，0.05 s 窗口包含约 1280 个采样点。三刃刀在 3500-5500 rpm 下刀齿通过频率约为 175-275 Hz，因此每个窗口大约包含 9-14 次刀齿通过，既能反映局部稳定切削能量，又不会把抬刀/停转造成的短时大冲击完全平均掉。",
        "当前窗口长度为 0.05 s。RMS 用于描述局部能量，峰峰值用于捕捉局部最大波动范围；二者后续进入 MAD 稳健阈值判断。",
    )

    add_step(
        doc,
        "3. MAD 稳健阈值判断异常冲击",
        "从窗口 RMS 和峰峰值序列中识别离群窗口，降低极端冲击对阈值估计的干扰。",
        "MAD（Median Absolute Deviation，中位数绝对偏差）是一种稳健离群检测方法。相比均值加标准差阈值，MAD 不容易被少量大幅异常点拉高阈值，因此更适合含有抬刀、停转、碰撞或偶发冲击的振动数据。",
        "你的 27-39 等单独获取的数据中存在抬刀与主轴停转动作，且这些动作会产生特别大的波动。若使用普通均值/标准差，异常冲击会抬高整体阈值，反而削弱异常识别能力；MAD 更符合这类数据的实际情况。",
        "当前阈值系数 k=8.0。该值偏保守，目的是只剔除明显异常冲击，尽量避免误删稳定切削中的正常振动波动。",
    )

    add_step(
        doc,
        "4. 异常窗口前后扩展 0.2 s 剔除",
        "删除异常尖峰附近的过渡影响，而不只是删除尖峰所在窗口。",
        "机械系统中的冲击或状态转换通常不是单点事件。抬刀、停转、切入切出和接触状态变化会在异常峰值前后形成短暂过渡响应。如果只删除峰值窗口，邻近窗口仍可能包含非稳定切削信息。",
        "你的数据已经截取为有效切削信号，但部分文件仍可能残留抬刀/停转或过渡段。将异常窗口前后各扩展 0.2 s，可以更稳妥地去除过渡影响，减少模型学习到非切削状态的风险。",
        "扩展长度为 0.2 s，对应约 5120 个采样点。该参数主要针对明显冲击附近的过渡段，若后续发现误删过多，可在敏感性分析中比较 0.1 s、0.2 s、0.3 s。",
    )

    add_step(
        doc,
        "5. 20-10000 Hz 带通滤波",
        "保留与切削振动相关的主要频率成分，削弱低频漂移和接近采样上限的高频噪声。",
        "铣削振动主要围绕主轴转频、刀齿通过频率、倍频、结构模态及冲击成分展开。极低频成分常与基线漂移、装夹缓慢变化、机床大尺度运动或传感器零漂有关；过高频成分更容易受到采集噪声和奈奎斯特边缘效应影响。带通滤波是振动监测和加工状态分析中常用的频域预处理方法。",
        "采样率为 25600 Hz，奈奎斯特频率为 12800 Hz。滤波上限设为 10000 Hz，可以保留大部分切削振动和高频冲击信息，同时避开接近 12800 Hz 的边缘噪声区域。滤波下限设为 20 Hz，不会删掉 58.3-91.7 Hz 的主轴转频，也不会删掉 175-275 Hz 的刀齿通过频率及其倍频，但可以去除更低频的漂移和缓慢趋势。",
        "当前采用 FFT 方式进行 20-10000 Hz 带通。该频带适合当前采样率和转速范围；若后续更换采样率或转速范围，应重新检查上限、下限与奈奎斯特频率、主轴转频、刀齿通过频率的关系。",
    )

    add_step(
        doc,
        "6. 重建连续时间轴",
        "在剔除异常片段后，将剩余有效信号整理为连续序列，便于后续切分、频域分析和模型输入。",
        "异常窗口被删除后，原始时间轴会出现空洞。对机器学习建模而言，片段输入通常需要连续、规则采样的序列；重建时间轴可以避免后续切分和频域计算受到时间索引断裂影响。",
        "你的目标是把每条稳定切削信号切分为一到三个同标签样本。若保留断裂时间轴，三等分和 CNN 输入会变得不规范；重建连续时间轴可以使每个输出片段具有统一采样间隔和连续序列形式。",
        "重建时间轴只改变索引表达，不改变保留下来的振动幅值序列。论文表述中应说明该步骤发生在异常片段剔除之后，目的是形成规则采样输入。",
    )

    add_step(
        doc,
        "7. 预处理后再切分",
        "保证切分得到的训练片段来自清理后的稳定切削信号，降低同一实验内部片段质量差异。",
        "信号级预处理应先于数据增强式切分完成，因为去偏置、异常段剔除和带通滤波作用于完整时间序列。若先切分再清理，异常冲击可能集中落入某个片段，使同标签片段质量不一致。",
        "当前数据中 212 条有效信号预处理后得到 586 个片段。预处理后比理论三等分数量少，是因为少数信号异常剔除后时长变短，按时长规则从三段降为两段或一段。这比强行三等分更稳妥。",
        "切分规则为：稳定切削时长 >=3 s 时三等分，2-3 s 时二等分，1-2 s 时保留单片段，<1 s 剔除。模型训练/验证/测试划分时，应按执行顺序分组，避免同一原始实验的多个片段进入不同数据集。",
    )

    doc.add_heading("三、建议写入论文的方法表述", level=1)
    add_para(
        doc,
        "为降低采集偏置、非稳定切削冲击和无关频率成分对模型训练的影响，本文首先对已截取的有效切削振动信号进行信号级预处理。具体包括：对各通道进行中位数去直流偏置；基于 0.05 s 滑动窗口计算 RMS 与峰峰值，并采用 MAD 稳健阈值识别异常冲击窗口；对异常窗口前后扩展 0.2 s 后剔除；随后在 20-10000 Hz 范围内进行带通滤波，并重建连续时间轴。预处理完成后，再依据有效切削时长将信号切分为一到三个同标签片段。训练、验证和测试划分按原始执行顺序分组完成，以避免同源片段造成数据泄漏。",
    )

    doc.add_heading("四、参考依据", level=1)
    refs = [
        ("Pontuale et al., Dynamical characterization of machining systems, arXiv:physics/0312148", "https://arxiv.org/abs/physics/0312148"),
        ("Machining vibrations", "https://en.wikipedia.org/wiki/Machining_vibrations"),
        ("Condition monitoring", "https://en.wikipedia.org/wiki/Condition_monitoring"),
        ("Hampel test / MAD-based outlier detection", "https://en.wikipedia.org/wiki/Hampel_identifier"),
    ]
    for text, url in refs:
        p = doc.add_paragraph(style=None)
        add_hyperlink(p, text, url)

    doc.add_heading("五、参数汇总", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "当前取值"
    hdr[2].text = "依据"
    for cell in hdr:
        set_cell_shading(cell, "F2F4F7")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    rows = [
        ("采样率", "25600 Hz", "来自当前预处理配置"),
        ("异常检测窗口", "0.05 s，约 1280 点", "可覆盖约 9-14 次刀齿通过"),
        ("MAD 阈值系数", "k = 8.0", "保守识别明显异常冲击"),
        ("异常扩展剔除", "前后各 0.2 s", "去除冲击附近过渡响应"),
        ("带通频带", "20-10000 Hz", "保留主轴转频、刀齿通过频率、倍频和主要切削振动；避开低频漂移和奈奎斯特边缘噪声"),
        ("切分规则", ">=3 s 三段；2-3 s 两段；1-2 s 一段；<1 s 剔除", "兼顾样本扩充和最小时长约束"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_width(table, [3.2, 4.0, 9.2])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("机床铣削振动信号预处理说明").font.size = Pt(9)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
