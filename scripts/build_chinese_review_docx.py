from __future__ import annotations

import re
import subprocess
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/paper/2026-09-14-sgrpn-manuscript-sections-1-9-zh.md"
OUTPUT = ROOT / "docs/paper/2026-09-17-sgrpn-chinese-review.docx"
CACHE = ROOT / ".cache/chinese_review_docx"
EQUATION_CACHE = CACHE / "equations"
MATH_RENDERER = ROOT / "scripts/render_math_png.py"
MATH_PYTHON = Path(r"D:\CodexPython\python.exe")


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5, bold=None, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


FONT_COMMANDS = (
    r"mathcal|mathrm|mathbf|mathbb|mathsf|mathtt|mathit|mathfrak"
    r"|mathnormal|boldsymbol|text|textrm|textbf|textit|operatorname"
)


def preprocess_latex(latex: str) -> str:
    # 1) 折叠全部空白（含换行）。mathtext 无法解析含换行的公式，且不会报错，
    #    而是静默把 LaTeX 源码当普通文字画出来——这正是审阅稿公式乱码的直接原因。
    #    源 Markdown 里同一公式常跨多行书写，因此这一步不可省略。
    latex = " ".join(latex.split())
    latex = latex.rstrip(".,，。")
    latex = re.sub(r"\\operatorname\{([^{}]+)\}", r"\\mathrm{\1}", latex)
    latex = re.sub(r"\\text\{([^{}]+)\}", r"\\mathrm{\1}", latex)
    # 2) mathtext 要求字体类命令必须带花括号：`\mathcal L` 非法，`\mathcal{L}` 合法。
    #    未加括号会直接抛 ParseFatalException。
    latex = re.sub(rf"(\\(?:{FONT_COMMANDS}))\s+([A-Za-z0-9]+)", r"\1{\2}", latex)
    return latex


def render_equation(latex: str, display: bool) -> Path:
    import hashlib

    cleaned = preprocess_latex(latex)
    digest = hashlib.sha256((("display:" if display else "inline:") + cleaned).encode("utf-8")).hexdigest()[:20]
    out = EQUATION_CACHE / f"{digest}.png"
    if out.exists():
        return out
    EQUATION_CACHE.mkdir(parents=True, exist_ok=True)
    command = [
        str(MATH_PYTHON),
        str(MATH_RENDERER),
        "--latex",
        cleaned,
        "--output",
        str(out),
    ]
    if display:
        command.append("--display")
    subprocess.run(command, check=True)
    return out


def latex_to_plain(latex: str) -> str:
    text = preprocess_latex(latex)
    replacements = {
        r"\mu": "μ", r"\sigma": "σ", r"\alpha": "α", r"\delta": "δ",
        r"\ell": "ℓ", r"\times": "×", r"\rightarrow": "→", r"\leq": "≤",
        r"\in": "∈", r"\ldots": "…", r"\cdot": "·", r"\qquad": "  ",
        r"\,": " ", "~": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\\(?:mathrm|mathcal)\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\(?:bar|hat)\s*([A-Za-z])", r"\1", text)
    text = re.sub(r"\\(?:bar|hat)\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\^\{?2\}?", "²", text)
    text = re.sub(r"\^\{?3\}?", "³", text)
    text = re.sub(r"_\{([^{}]+)\}", r"_\1", text)
    text = re.sub(r"\^\{([^{}]+)\}", r"^\1", text)
    text = text.replace(r"\(", "(").replace(r"\)", ")")
    text = text.replace(r"\[", "[").replace(r"\]", "]")
    text = text.replace(r"\{", "{").replace(r"\}", "}")
    text = re.sub(r"\\[A-Za-z]+", "", text)
    return text


INLINE_TOKEN = re.compile(r"(\\\(.+?\\\)|\*\*.+?\*\*|\*[^*\n]+?\*|\[[^\]]+\]\([^)]+\))")


def add_inline_content(paragraph, text: str, font_size=10.5, allow_equation_images=True):
    pos = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=font_size)
        token = match.group(0)
        if token.startswith(r"\("):
            latex = token[2:-2]
            if allow_equation_images:
                image_path = render_equation(latex, display=False)
                with Image.open(image_path) as im:
                    aspect = im.width / max(im.height, 1)
                height_pt = min(12.5, 390 / max(aspect, 0.1))
                paragraph.add_run().add_picture(str(image_path), height=Pt(height_pt))
            else:
                run = paragraph.add_run(latex_to_plain(latex))
                set_run_font(run, east_asia="Cambria Math", latin="Cambria Math", size=font_size)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=font_size, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=font_size, italic=True)
        else:
            label = token[1:token.index("]")]
            run = paragraph.add_run(label)
            set_run_font(run, size=font_size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=font_size)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.first_line_indent = Pt(21)

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(14)

    for style_name, size, before, after in (
        ("Heading 1", 14, 14, 6), ("Heading 2", 12, 10, 4), ("Heading 3", 11, 8, 3)
    ):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)

    for style_name in ("List Number", "List Bullet"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10.5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    set_run_font(run, size=9)


def add_display_equation(doc: Document, latex: str):
    path = render_equation(latex, display=True)
    with Image.open(path) as im:
        aspect = im.width / max(im.height, 1)
    width_in = min(6.2, 0.72 * aspect)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(width_in))


def parse_table(lines: list[str]) -> list[list[str]]:
    return [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def add_table(doc: Document, lines: list[str]):
    rows = parse_table(lines)
    if len(rows) >= 2 and is_separator_row(rows[1]):
        rows.pop(1)
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.05
            if r_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade_cell(cell, "1F4E78")
            elif r_idx % 2 == 0:
                shade_cell(cell, "EAF2F8")
            value = row[c_idx] if c_idx < len(row) else ""
            add_inline_content(p, value, font_size=7.4 if cols >= 8 else 8.2, allow_equation_images=False)
            for run in p.runs:
                if r_idx == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            if c_idx > 0 and len(value) < 28:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    after.paragraph_format.first_line_indent = Pt(0)


def add_image(doc: Document, md_path: str):
    image_path = (SOURCE.parent / md_path).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    if not image_path.exists():
        run = p.add_run(f"[图片缺失：{md_path}]")
        set_run_font(run, size=10)
        return
    with Image.open(image_path) as im:
        width_px, height_px = im.size
    aspect = width_px / max(height_px, 1)
    width_in = min(6.15, 5.6 * aspect)
    p.add_run().add_picture(str(image_path), width=Inches(width_in))


def add_body_paragraph(doc: Document, text: str, reference_mode=False):
    style = None
    if re.match(r"^\d+\.\s+", text) and not reference_mode:
        style = "List Number"
        text = re.sub(r"^\d+\.\s+", "", text)
    p = doc.add_paragraph(style=style)
    if reference_mode:
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.hanging_indent = Pt(18)
        p.paragraph_format.line_spacing = 1.1
        p.paragraph_format.space_after = Pt(4)
        size = 9
    else:
        size = 10.5
    add_inline_content(p, text, font_size=size)
    return p


def build_document():
    EQUATION_CACHE.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    add_page_number(section.footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    in_references = False
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith(r"\["):
            equation_lines = []
            first = stripped[2:].strip()
            if first:
                equation_lines.append(first)
            i += 1
            while i < len(lines) and not lines[i].strip().endswith(r"\]"):
                equation_lines.append(lines[i].strip())
                i += 1
            if i < len(lines):
                last = lines[i].strip()
                if last != r"\]":
                    equation_lines.append(last[:-2].strip())
            add_display_equation(doc, "\n".join(equation_lines))
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            add_image(doc, image_match.group(2))
            i += 1
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline_content(p, stripped[2:].strip(), font_size=18)
            i += 1
            continue
        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            doc.add_paragraph(heading, style="Heading 1")
            in_references = heading == "参考文献"
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_paragraph(stripped[4:].strip(), style="Heading 2")
            i += 1
            continue
        if stripped.startswith("#### "):
            doc.add_paragraph(stripped[5:].strip(), style="Heading 3")
            i += 1
            continue
        if stripped.startswith("*图") and stripped.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[1:-1])
            set_run_font(run, size=9)
            i += 1
            continue
        if stripped.startswith("**表") and stripped.endswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:-2])
            set_run_font(run, size=9, bold=True)
            i += 1
            continue
        if stripped.startswith("**面板") and stripped.endswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:-2])
            set_run_font(run, size=9.5, bold=True)
            i += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline_content(p, stripped[2:], font_size=9.5)
            i += 1
            continue
        add_body_paragraph(doc, stripped, reference_mode=in_references)
        i += 1

    core = doc.core_properties
    core.title = "面向铣削表面粗糙度的安全门控残差神经融合与组拆分保形不确定性量化"
    core.subject = "中文审阅稿"
    core.author = ""
    core.keywords = "表面粗糙度预测；铣削；选择性传感融合；负迁移；组保形预测；不确定性量化"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
